import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from typing import Dict, List
import pandas as pd
from config import (
    BAIN_RED, BAIN_DARK_GREY, BAIN_MEDIUM_GREY, BAIN_LIGHT_GREY,
    BAIN_WHITE, BAIN_BORDER_GREY,
    TITLE_FONT, TITLE_SIZE, TITLE_BOLD,
    HEADER_FONT, HEADER_SIZE, HEADER_BOLD,
    DATA_FONT, DATA_SIZE,
    CATEGORY_FONT, CATEGORY_SIZE, CATEGORY_ITALIC, CATEGORY_COLOR,
    BLOCK_GAP_ROWS
)

def hex_to_argb(hex_color):
    hex_color = hex_color.replace('#', '')
    return 'FF' + hex_color

class FormulaExcelFormatter:
    """Creates formula-driven Excel output with Raw_Data sheet and TDW_Cuts sheet."""
    
    def __init__(self, include_percentages=False):
        self.wb = openpyxl.Workbook()
        self.include_percentages = include_percentages
        self.current_row = 1
        
        # Column mapping: column_name -> column_letter
        self.column_map = {}
        # Data range tracking
        self.data_start_row = 2  # Row 1 is header
        self.data_end_row = None
        
        # Remove default sheet
        if 'Sheet' in self.wb.sheetnames:
            self.wb.remove(self.wb['Sheet'])
        
        # Create sheets
        self.ws_raw = self.wb.create_sheet('Raw_Data', 0)
        self.ws_cuts = self.wb.create_sheet('TDW_Cuts', 1)
        
        # Setup styles
        self._setup_styles()
    
    def _setup_styles(self):
        """Setup reusable cell styles."""
        self.title_font = Font(
            name=TITLE_FONT, size=TITLE_SIZE, bold=TITLE_BOLD,
            color=hex_to_argb(BAIN_DARK_GREY)
        )
        self.header_font = Font(
            name=HEADER_FONT, size=HEADER_SIZE, bold=HEADER_BOLD,
            color=hex_to_argb(BAIN_WHITE)
        )
        self.header_fill = PatternFill(
            start_color=hex_to_argb(BAIN_DARK_GREY),
            end_color=hex_to_argb(BAIN_DARK_GREY),
            fill_type='solid'
        )
        self.data_font = Font(
            name=DATA_FONT, size=DATA_SIZE,
            color=hex_to_argb(BAIN_DARK_GREY)
        )
        self.category_font = Font(
            name=CATEGORY_FONT, size=CATEGORY_SIZE,
            italic=CATEGORY_ITALIC,
            color=hex_to_argb(CATEGORY_COLOR)
        )
        self.thin_border = Border(
            left=Side(style='thin', color=hex_to_argb(BAIN_BORDER_GREY)),
            right=Side(style='thin', color=hex_to_argb(BAIN_BORDER_GREY)),
            top=Side(style='thin', color=hex_to_argb(BAIN_BORDER_GREY)),
            bottom=Side(style='thin', color=hex_to_argb(BAIN_BORDER_GREY))
        )
        self.center_align = Alignment(horizontal='center', vertical='center')
        self.left_align = Alignment(horizontal='left', vertical='center')
    
    def add_raw_data(self, df: pd.DataFrame):
        """Add raw data with all columns and track column positions."""
        # Write headers and build column map
        for col_idx, col_name in enumerate(df.columns, 1):
            col_letter = get_column_letter(col_idx)
            self.column_map[str(col_name)] = col_letter
            
            cell = self.ws_raw.cell(row=1, column=col_idx, value=col_name)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border
        
        # Write data
        for row_idx, row in enumerate(df.itertuples(index=False), 2):
            for col_idx, value in enumerate(row, 1):
                cell = self.ws_raw.cell(row=row_idx, column=col_idx)
                if pd.isna(value):
                    cell.value = None
                else:
                    cell.value = value
                cell.font = self.data_font
                cell.border = self.thin_border
                cell.alignment = self.left_align
        
        self.data_end_row = len(df) + 1
        
        # Create Excel Table for formatting (still useful for Excel users)
        table_ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
        table = Table(displayName='tblRawData', ref=table_ref)
        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        table.tableStyleInfo = style
        self.ws_raw.add_table(table)
        
        # Auto-adjust column widths
        for col_idx, col_name in enumerate(df.columns, 1):
            max_length = max(
                len(str(col_name)),
                df[col_name].astype(str).str.len().max() if len(df) > 0 else 0
            )
            self.ws_raw.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 50)
    
    def _get_col_range(self, col_name: str) -> str:
        """Get A1-style column range for a column."""
        col_letter = self.column_map.get(str(col_name))
        if not col_letter:
            # Fallback: try to find by case-insensitive match
            for key, val in self.column_map.items():
                if key.lower() == str(col_name).lower():
                    col_letter = val
                    break
        if not col_letter:
            raise ValueError(f"Column '{col_name}' not found in raw data")
        return f"'Raw_Data'!${col_letter}${self.data_start_row}:${col_letter}${self.data_end_row}"
    
    def add_cuts(self, cuts: Dict, demo_columns: List[str]):
        """Add all 8 cuts with formulas."""
        self.ws_cuts.column_dimensions['A'].width = 3
        
        # Cut 1-3: Demographic summaries
        for i, demo_col in enumerate(demo_columns, 1):
            cut_key = f'cut{i}'
            if cut_key in cuts:
                self._add_demographic_cut_formulas(
                    cuts[cut_key], i, demo_col
                )
        
        # Cut 4: L1 Split
        if 'cut4' in cuts:
            self._add_l1_split_formulas(cuts['cut4'])
        
        # Cut 5: L2 vs L1
        if 'cut5' in cuts:
            self._add_l2_vs_l1_formulas(cuts['cut5'])
        
        # Cut 6-8: L1 vs Demographics
        for i, demo_col in enumerate(demo_columns, 6):
            cut_key = f'cut{i}'
            if cut_key in cuts:
                self._add_l1_vs_demo_formulas(
                    cuts[cut_key], i, demo_col
                )
    
    def _add_demographic_cut_formulas(self, cut_data: Dict, cut_num: int, demo_col: str):
        """Add demographic summary with COUNTIF formulas."""
        start_row = self.current_row
        
        # Title
        title_cell = self.ws_cuts.cell(row=start_row, column=2, value=cut_data['title'])
        title_cell.font = self.title_font
        self.current_row += 1
        
        # Headers
        self.ws_cuts.cell(row=self.current_row, column=2, value=demo_col).font = self.header_font
        self.ws_cuts.cell(row=self.current_row, column=2).fill = self.header_fill
        self.ws_cuts.cell(row=self.current_row, column=2).alignment = self.center_align
        self.ws_cuts.cell(row=self.current_row, column=2).border = self.thin_border
        
        self.ws_cuts.cell(row=self.current_row, column=3, value='Count').font = self.header_font
        self.ws_cuts.cell(row=self.current_row, column=3).fill = self.header_fill
        self.ws_cuts.cell(row=self.current_row, column=3).alignment = self.center_align
        self.ws_cuts.cell(row=self.current_row, column=3).border = self.thin_border
        
        if self.include_percentages:
            self.ws_cuts.cell(row=self.current_row, column=4, value='%').font = self.header_font
            self.ws_cuts.cell(row=self.current_row, column=4).fill = self.header_fill
            self.ws_cuts.cell(row=self.current_row, column=4).alignment = self.center_align
            self.ws_cuts.cell(row=self.current_row, column=4).border = self.thin_border
        
        self.current_row += 1
        data_start_row = self.current_row
        
        # Get column range for formulas
        col_range = self._get_col_range(demo_col)
        
        # Data rows with formulas
        categories = cut_data['categories']
        for i, cat in enumerate(categories):
            row = self.current_row
            
            # Category name (grey italic)
            cat_cell = self.ws_cuts.cell(row=row, column=2, value=cat)
            cat_cell.font = self.category_font
            cat_cell.border = self.thin_border
            cat_cell.alignment = self.left_align
            
            # COUNTIF formula
            count_cell = self.ws_cuts.cell(row=row, column=3)
            if cat == '(Blank)':
                count_cell.value = f'=COUNTBLANK({col_range})'
            else:
                count_cell.value = f'=COUNTIF({col_range},B{row})'
            count_cell.font = self.data_font
            count_cell.border = self.thin_border
            count_cell.alignment = self.center_align
            
            # Percentage formula
            if self.include_percentages:
                pct_cell = self.ws_cuts.cell(row=row, column=4)
                pct_cell.value = f'=C{row}/C{data_start_row + len(categories)}'
                pct_cell.number_format = '0.0%'
                pct_cell.font = self.data_font
                pct_cell.border = self.thin_border
                pct_cell.alignment = self.center_align
            
            self.current_row += 1
        
        # Total row
        total_row = self.current_row
        total_cell = self.ws_cuts.cell(row=total_row, column=2, value='Total')
        total_cell.font = Font(name=DATA_FONT, size=DATA_SIZE, bold=True, color=hex_to_argb(BAIN_DARK_GREY))
        total_cell.border = self.thin_border
        
        total_count = self.ws_cuts.cell(row=total_row, column=3)
        total_count.value = f'=SUM(C{data_start_row}:C{total_row - 1})'
        total_count.font = Font(name=DATA_FONT, size=DATA_SIZE, bold=True, color=hex_to_argb(BAIN_DARK_GREY))
        total_count.border = self.thin_border
        total_count.alignment = self.center_align
        
        if self.include_percentages:
            total_pct = self.ws_cuts.cell(row=total_row, column=4, value=1.0)
            total_pct.number_format = '0.0%'
            total_pct.font = Font(name=DATA_FONT, size=DATA_SIZE, bold=True, color=hex_to_argb(BAIN_DARK_GREY))
            total_pct.border = self.thin_border
            total_pct.alignment = self.center_align
        
        self.current_row += 1 + BLOCK_GAP_ROWS
        
        # Set column widths
        self.ws_cuts.column_dimensions['B'].width = 40
        self.ws_cuts.column_dimensions['C'].width = 15
        if self.include_percentages:
            self.ws_cuts.column_dimensions['D'].width = 15
    
    def _add_l1_split_formulas(self, cut_data: Dict):
        """Add L1 split with COUNTIF formulas."""
        start_row = self.current_row
        
        title_cell = self.ws_cuts.cell(row=start_row, column=2, value=cut_data['title'])
        title_cell.font = self.title_font
        self.current_row += 1
        
        # Headers
        self.ws_cuts.cell(row=self.current_row, column=2, value='L1 Category').font = self.header_font
        self.ws_cuts.cell(row=self.current_row, column=2).fill = self.header_fill
        self.ws_cuts.cell(row=self.current_row, column=2).alignment = self.center_align
        self.ws_cuts.cell(row=self.current_row, column=2).border = self.thin_border
        
        self.ws_cuts.cell(row=self.current_row, column=3, value='Count').font = self.header_font
        self.ws_cuts.cell(row=self.current_row, column=3).fill = self.header_fill
        self.ws_cuts.cell(row=self.current_row, column=3).alignment = self.center_align
        self.ws_cuts.cell(row=self.current_row, column=3).border = self.thin_border
        
        if self.include_percentages:
            self.ws_cuts.cell(row=self.current_row, column=4, value='%').font = self.header_font
            self.ws_cuts.cell(row=self.current_row, column=4).fill = self.header_fill
            self.ws_cuts.cell(row=self.current_row, column=4).alignment = self.center_align
            self.ws_cuts.cell(row=self.current_row, column=4).border = self.thin_border
        
        self.current_row += 1
        data_start_row = self.current_row
        
        # Get L1_Parsed column range
        l1_range = self._get_col_range('L1_Parsed')
        
        categories = cut_data['categories']
        for i, cat in enumerate(categories):
            row = self.current_row
            
            cat_cell = self.ws_cuts.cell(row=row, column=2, value=cat)
            cat_cell.font = self.data_font
            cat_cell.border = self.thin_border
            cat_cell.alignment = self.left_align
            
            count_cell = self.ws_cuts.cell(row=row, column=3)
            count_cell.value = f'=COUNTIF({l1_range},"{cat}")'
            count_cell.font = self.data_font
            count_cell.border = self.thin_border
            count_cell.alignment = self.center_align
            
            if self.include_percentages:
                pct_cell = self.ws_cuts.cell(row=row, column=4)
                pct_cell.value = f'=C{row}/C{data_start_row + len(categories)}'
                pct_cell.number_format = '0.0%'
                pct_cell.font = self.data_font
                pct_cell.border = self.thin_border
                pct_cell.alignment = self.center_align
            
            self.current_row += 1
        
        # Total
        total_row = self.current_row
        self.ws_cuts.cell(row=total_row, column=2, value='Total').font = Font(
            name=DATA_FONT, size=DATA_SIZE, bold=True, color=hex_to_argb(BAIN_DARK_GREY)
        )
        self.ws_cuts.cell(row=total_row, column=2).border = self.thin_border
        
        total_count = self.ws_cuts.cell(row=total_row, column=3)
        total_count.value = f'=SUM(C{data_start_row}:C{total_row - 1})'
        total_count.font = Font(name=DATA_FONT, size=DATA_SIZE, bold=True, color=hex_to_argb(BAIN_DARK_GREY))
        total_count.border = self.thin_border
        total_count.alignment = self.center_align
        
        if self.include_percentages:
            total_pct = self.ws_cuts.cell(row=total_row, column=4, value=1.0)
            total_pct.number_format = '0.0%'
            total_pct.font = Font(name=DATA_FONT, size=DATA_SIZE, bold=True, color=hex_to_argb(BAIN_DARK_GREY))
            total_pct.border = self.thin_border
            total_pct.alignment = self.center_align
        
        self.current_row += 1 + BLOCK_GAP_ROWS
    
    def _add_l2_vs_l1_formulas(self, cut_data: Dict):
        """Add L2 vs L1 cross cut with COUNTIFS formulas."""
        start_row = self.current_row
        
        title_cell = self.ws_cuts.cell(row=start_row, column=2, value=cut_data['title'])
        title_cell.font = self.title_font
        self.current_row += 1
        
        # Headers
        headers = [''] + cut_data['columns']
        for col_idx, header in enumerate(headers, 2):
            cell = self.ws_cuts.cell(row=self.current_row, column=col_idx, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border
        
        header_row = self.current_row
        self.current_row += 1
        
        # Get column ranges
        l2_full_range = self._get_col_range('L2_Full')
        l1_parsed_range = self._get_col_range('L1_Parsed')
        
        rows = cut_data['rows']
        for i, row_name in enumerate(rows):
            row = self.current_row
            
            row_header = self.ws_cuts.cell(row=row, column=2, value=row_name)
            row_header.font = self.header_font
            row_header.fill = self.header_fill
            row_header.alignment = self.left_align
            row_header.border = self.thin_border
            
            for col_idx, col_name in enumerate(cut_data['columns'], 3):
                cell = self.ws_cuts.cell(row=row, column=col_idx)
                col_letter = get_column_letter(col_idx)
                cell.value = f'=COUNTIFS({l2_full_range},$B{row},{l1_parsed_range},{col_letter}${header_row})'
                cell.font = self.data_font
                cell.border = self.thin_border
                cell.alignment = self.center_align
            
            self.current_row += 1
        
        self.current_row += BLOCK_GAP_ROWS
        
        # Set column widths
        self.ws_cuts.column_dimensions['B'].width = 35
        for col_idx in range(3, 3 + len(cut_data['columns'])):
            self.ws_cuts.column_dimensions[get_column_letter(col_idx)].width = 15
    
    def _add_l1_vs_demo_formulas(self, cut_data: Dict, cut_num: int, demo_col: str):
        """Add L1 vs Demographic cross cut with COUNTIFS formulas."""
        start_row = self.current_row
        
        title_cell = self.ws_cuts.cell(row=start_row, column=2, value=cut_data['title'])
        title_cell.font = self.title_font
        self.current_row += 1
        
        # Headers
        headers = [''] + cut_data['columns']
        for col_idx, header in enumerate(headers, 2):
            cell = self.ws_cuts.cell(row=self.current_row, column=col_idx, value=header)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_align
            cell.border = self.thin_border
        
        header_row = self.current_row
        self.current_row += 1
        
        # Get column ranges
        l1_parsed_range = self._get_col_range('L1_Parsed')
        demo_range = self._get_col_range(demo_col)
        
        rows = cut_data['rows']
        columns = cut_data['columns']
        
        for i, row_name in enumerate(rows):
            row = self.current_row
            
            row_header = self.ws_cuts.cell(row=row, column=2, value=row_name)
            row_header.font = self.header_font
            row_header.fill = self.header_fill
            row_header.alignment = self.left_align
            row_header.border = self.thin_border
            
            for col_idx, col_name in enumerate(columns, 3):
                cell = self.ws_cuts.cell(row=row, column=col_idx)
                col_letter = get_column_letter(col_idx)
                
                if col_name == '(Blank)':
                    cell.value = f'=COUNTIFS({l1_parsed_range},$B{row},{demo_range},"")'
                else:
                    cell.value = f'=COUNTIFS({l1_parsed_range},$B{row},{demo_range},{col_letter}${header_row})'
                
                cell.font = self.data_font
                cell.border = self.thin_border
                cell.alignment = self.center_align
            
            self.current_row += 1
        
        self.current_row += BLOCK_GAP_ROWS
        
        # Set column widths
        self.ws_cuts.column_dimensions['B'].width = 20
        for col_idx in range(3, 3 + len(columns)):
            self.ws_cuts.column_dimensions[get_column_letter(col_idx)].width = 18
    
    def save(self, filepath: str):
        """Save the workbook."""
        # Freeze panes for better viewing
        self.ws_cuts.freeze_panes = 'B2'
        
        self.wb.save(filepath)
        return filepath